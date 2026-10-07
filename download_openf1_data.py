from pathlib import Path
import time

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIG
# ============================================================

YEAR = 2026
COUNTRY = "Japan"
SESSION_NAME = "Race"

OUT_DIR = Path("japan_2026_openf1")
OUT_DIR.mkdir(parents=True, exist_ok=True)

BASE = "https://api.openf1.org/v1"

http = requests.Session()
http.headers.update(
    {
        "User-Agent": "F1-research-data/1.0",
        "Accept": "application/json",
    }
)


# ============================================================
# ROBUST API GET
# ============================================================

def get_json(endpoint, params=None, tries=5, timeout=60):
    url = f"{BASE}/{endpoint}"

    for attempt in range(1, tries + 1):
        try:
            r = http.get(url, params=params, timeout=timeout)

            if r.status_code == 429:
                wait = min(5 * attempt, 30)
                print(f"Rate limited. Waiting {wait}s...")
                time.sleep(wait)
                continue

            r.raise_for_status()
            return r.json()

        except requests.RequestException as exc:
            if attempt == tries:
                raise RuntimeError(
                    f"Failed GET {url} with params={params}\n{exc}"
                ) from exc

            wait = 2 * attempt
            print(
                f"Request failed ({attempt}/{tries}): {exc}\n"
                f"Retrying in {wait}s..."
            )
            time.sleep(wait)

    raise RuntimeError(f"Failed GET {url}")


# ============================================================
# EXCEL EXPORT
# ============================================================

def save_formatted_excel(df, path, sheet_name, table_name):
    """
    Save a DataFrame to a readable .xlsx workbook.
    """

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False,
            na_rep="",
        )

        ws = writer.book[sheet_name]

        # Keep header and first four descriptive columns visible.
        ws.freeze_panes = "E2"
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90

        # Header formatting.
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        ws.row_dimensions[1].height = 30

        # Add a real Excel table with filters and striped rows.
        if ws.max_row >= 2 and ws.max_column >= 1:
            last_col = get_column_letter(ws.max_column)

            table = Table(
                displayName=table_name,
                ref=f"A1:{last_col}{ws.max_row}",
            )

            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            ws.add_table(table)

        preferred_widths = {
            "Driver": 12,
            "DriverFullName": 24,
            "DriverNumber": 14,
            "Team": 22,
            "LapNumber": 12,
            "LapTime_seconds": 18,
            "Sector1Time_seconds": 20,
            "Sector2Time_seconds": 20,
            "Sector3Time_seconds": 20,
            "SectorSum_seconds": 18,
            "SectorSumMinusLap_seconds": 24,
            "Stint": 10,
            "Compound": 13,
            "TyreAgeAtStintStart": 22,
            "EstimatedTyreAgeAtLapStart": 28,
            "IsPitOutLap": 14,
            "PitLaneDuration_seconds": 24,
            "StationaryStopDuration_seconds": 29,
            "I1Speed_kmh": 15,
            "I2Speed_kmh": 15,
            "SpeedTrap_kmh": 17,
            "LapStartUTC": 28,
        }

        # Column widths.
        for col_idx, col_name in enumerate(df.columns, start=1):
            letter = get_column_letter(col_idx)

            if col_name in preferred_widths:
                width = preferred_widths[col_name]

            else:
                sample = df[col_name].astype(str).head(500)

                values = [
                    len(v)
                    for v in sample
                    if v.lower() not in {"nan", "nat", "<na>"}
                ]

                max_len = max(
                    [len(str(col_name))] + values
                )

                width = min(
                    max(max_len + 2, 14),
                    35,
                )

            ws.column_dimensions[letter].width = width

        # Row height and alignment.
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 20

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=False,
                )

        column_positions = {
            name: idx + 1
            for idx, name in enumerate(df.columns)
        }

        three_decimal_columns = {
            "LapTime_seconds",
            "Sector1Time_seconds",
            "Sector2Time_seconds",
            "Sector3Time_seconds",
            "SectorSum_seconds",
            "SectorSumMinusLap_seconds",
            "PitLaneDuration_seconds",
            "StationaryStopDuration_seconds",
        }

        one_decimal_columns = {
            "I1Speed_kmh",
            "I2Speed_kmh",
            "SpeedTrap_kmh",
        }

        integer_columns = {
            "DriverNumber",
            "LapNumber",
            "Stint",
        }

        for name in three_decimal_columns:
            if name in column_positions:
                col_idx = column_positions[name]

                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(
                        row=row_idx,
                        column=col_idx,
                    ).number_format = "0.000"

        for name in one_decimal_columns:
            if name in column_positions:
                col_idx = column_positions[name]

                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(
                        row=row_idx,
                        column=col_idx,
                    ).number_format = "0.0"

        for name in integer_columns:
            if name in column_positions:
                col_idx = column_positions[name]

                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(
                        row=row_idx,
                        column=col_idx,
                    ).number_format = "0"


# ============================================================
# 1) FIND THE 2026 JAPANESE GP
# ============================================================

print("=" * 72)
print("Finding 2026 Japanese Grand Prix in OpenF1")
print("=" * 72)

meetings = get_json(
    "meetings",
    {
        "year": YEAR,
        "country_name": COUNTRY,
    },
)

if not meetings:
    raise RuntimeError(
        f"No OpenF1 meeting found for {COUNTRY}, {YEAR}."
    )

gp_candidates = [
    m
    for m in meetings
    if (
        "grand prix" in str(
            m.get("meeting_name", "")
        ).lower()
        or "grand prix" in str(
            m.get("meeting_official_name", "")
        ).lower()
    )
]

meeting = (
    gp_candidates[0]
    if gp_candidates
    else meetings[0]
)

meeting_key = meeting["meeting_key"]

print(f"Meeting: {meeting.get('meeting_name')}")
print(
    f"Official name: "
    f"{meeting.get('meeting_official_name')}"
)
print(f"Meeting key: {meeting_key}")


# ============================================================
# 2) FIND THE RACE SESSION
# ============================================================

sessions = get_json(
    "sessions",
    {
        "meeting_key": meeting_key,
        "session_name": SESSION_NAME,
    },
)

if not sessions:
    sessions = get_json(
        "sessions",
        {"meeting_key": meeting_key},
    )

    sessions = [
        s
        for s in sessions
        if (
            str(
                s.get("session_type", "")
            ).lower()
            == "race"
            or str(
                s.get("session_name", "")
            ).lower()
            == "race"
        )
    ]

if not sessions:
    raise RuntimeError(
        f"No Race session found "
        f"for meeting_key={meeting_key}."
    )

race_session = sessions[0]
session_key = race_session["session_key"]

print(f"Session: {race_session.get('session_name')}")
print(f"Session key: {session_key}")


# ============================================================
# 3) DOWNLOAD LAPS, DRIVERS, STINTS, PIT STOPS
# ============================================================

print("\nDownloading lap timing...")

laps_json = get_json(
    "laps",
    {"session_key": session_key},
)

if not laps_json:
    raise RuntimeError(
        f"OpenF1 returned no lap data "
        f"for session_key={session_key}."
    )

print("Downloading driver metadata...")

drivers_json = get_json(
    "drivers",
    {"session_key": session_key},
)

print("Downloading stint data...")

stints_json = get_json(
    "stints",
    {"session_key": session_key},
)

print("Downloading pit-stop data...")

pit_json = get_json(
    "pit",
    {"session_key": session_key},
)


# ============================================================
# 4) BUILD LAP TABLE
# ============================================================

laps = pd.DataFrame(laps_json)

rename_laps = {
    "driver_number": "DriverNumber",
    "lap_number": "LapNumber",
    "lap_duration": "LapTime_seconds",
    "duration_sector_1": "Sector1Time_seconds",
    "duration_sector_2": "Sector2Time_seconds",
    "duration_sector_3": "Sector3Time_seconds",
    "date_start": "LapStartUTC",
    "is_pit_out_lap": "IsPitOutLap",
    "i1_speed": "I1Speed_kmh",
    "i2_speed": "I2Speed_kmh",
    "st_speed": "SpeedTrap_kmh",
}

laps = laps.rename(
    columns={
        k: v
        for k, v in rename_laps.items()
        if k in laps.columns
    }
)

numeric_cols = [
    "DriverNumber",
    "LapNumber",
    "LapTime_seconds",
    "Sector1Time_seconds",
    "Sector2Time_seconds",
    "Sector3Time_seconds",
    "I1Speed_kmh",
    "I2Speed_kmh",
    "SpeedTrap_kmh",
]

for col in numeric_cols:
    if col in laps.columns:
        laps[col] = pd.to_numeric(
            laps[col],
            errors="coerce",
        )

if "LapNumber" in laps.columns:
    laps["LapNumber"] = laps[
        "LapNumber"
    ].astype("Int64")

if "DriverNumber" in laps.columns:
    laps["DriverNumber"] = laps[
        "DriverNumber"
    ].astype("Int64")


# ============================================================
# 5) MERGE DRIVER NAMES / TEAMS
# ============================================================

if drivers_json:
    drivers = pd.DataFrame(drivers_json)

    keep_driver = [
        c
        for c in [
            "driver_number",
            "name_acronym",
            "full_name",
            "team_name",
        ]
        if c in drivers.columns
    ]

    drivers = (
        drivers[keep_driver]
        .drop_duplicates(
            subset=["driver_number"]
        )
        .rename(
            columns={
                "driver_number": "DriverNumber",
                "name_acronym": "Driver",
                "full_name": "DriverFullName",
                "team_name": "Team",
            }
        )
    )

    drivers["DriverNumber"] = pd.to_numeric(
        drivers["DriverNumber"],
        errors="coerce",
    ).astype("Int64")

    laps = laps.merge(
        drivers,
        on="DriverNumber",
        how="left",
    )


# ============================================================
# 6) ADD STINT / COMPOUND INFORMATION
# ============================================================

laps["Stint"] = pd.NA
laps["Compound"] = pd.NA
laps["TyreAgeAtStintStart"] = pd.NA

if stints_json:
    stints = pd.DataFrame(stints_json)

    for _, s in stints.iterrows():
        dnum = s.get("driver_number")
        lap_start = s.get("lap_start")
        lap_end = s.get("lap_end")

        if (
            pd.isna(dnum)
            or pd.isna(lap_start)
        ):
            continue

        if pd.isna(lap_end):
            lap_end = laps.loc[
                laps["DriverNumber"]
                == int(dnum),
                "LapNumber",
            ].max()

        mask = (
            (
                laps["DriverNumber"]
                == int(dnum)
            )
            & (
                laps["LapNumber"]
                >= int(lap_start)
            )
            & (
                laps["LapNumber"]
                <= int(lap_end)
            )
        )

        laps.loc[
            mask,
            "Stint",
        ] = s.get("stint_number")

        laps.loc[
            mask,
            "Compound",
        ] = s.get("compound")

        laps.loc[
            mask,
            "TyreAgeAtStintStart",
        ] = s.get("tyre_age_at_start")


# ============================================================
# RECONSTRUCT TYRE AGE AT EACH LAP START
# ============================================================

laps["EstimatedTyreAgeAtLapStart"] = pd.NA

if stints_json:
    for _, s in stints.iterrows():

        dnum = s.get("driver_number")
        lap_start = s.get("lap_start")
        lap_end = s.get("lap_end")
        tyre_age = s.get("tyre_age_at_start")

        if (
            pd.isna(dnum)
            or pd.isna(lap_start)
            or pd.isna(tyre_age)
        ):
            continue

        if pd.isna(lap_end):
            lap_end = laps.loc[
                laps["DriverNumber"]
                == int(dnum),
                "LapNumber",
            ].max()

        mask = (
            (
                laps["DriverNumber"]
                == int(dnum)
            )
            & (
                laps["LapNumber"]
                >= int(lap_start)
            )
            & (
                laps["LapNumber"]
                <= int(lap_end)
            )
        )

        laps.loc[
            mask,
            "EstimatedTyreAgeAtLapStart",
        ] = (
            float(tyre_age)
            + (
                laps.loc[
                    mask,
                    "LapNumber",
                ].astype(float)
                - float(lap_start)
            )
        )


# ============================================================
# 7) ADD PIT-STOP INFORMATION
# ============================================================

laps["PitLaneDuration_seconds"] = pd.NA
laps["StationaryStopDuration_seconds"] = pd.NA

if pit_json:

    pits = pd.DataFrame(pit_json)

    for _, p in pits.iterrows():

        dnum = p.get("driver_number")
        lap_no = p.get("lap_number")

        if (
            pd.isna(dnum)
            or pd.isna(lap_no)
        ):
            continue

        mask = (
            (
                laps["DriverNumber"]
                == int(dnum)
            )
            & (
                laps["LapNumber"]
                == int(lap_no)
            )
        )

        lane_duration = p.get(
            "lane_duration",
            p.get("pit_duration"),
        )

        laps.loc[
            mask,
            "PitLaneDuration_seconds",
        ] = lane_duration

        laps.loc[
            mask,
            "StationaryStopDuration_seconds",
        ] = p.get("stop_duration")


# ============================================================
# 8) INTERNAL CONSISTENCY CHECK
# ============================================================

sector_cols = [
    "Sector1Time_seconds",
    "Sector2Time_seconds",
    "Sector3Time_seconds",
]

if all(
    c in laps.columns
    for c in sector_cols
):

    laps["SectorSum_seconds"] = (
        laps[sector_cols].sum(
            axis=1,
            min_count=3,
        )
    )

    if "LapTime_seconds" in laps.columns:

        laps[
            "SectorSumMinusLap_seconds"
        ] = (
            laps["SectorSum_seconds"]
            - laps["LapTime_seconds"]
        )


# ============================================================
# 9) CLEAN COLUMN ORDER
# ============================================================

preferred_order = [
    "Driver",
    "DriverFullName",
    "DriverNumber",
    "Team",
    "LapNumber",
    "LapTime_seconds",
    "Sector1Time_seconds",
    "Sector2Time_seconds",
    "Sector3Time_seconds",
    "SectorSum_seconds",
    "SectorSumMinusLap_seconds",
    "Stint",
    "Compound",
    "TyreAgeAtStintStart",
    "EstimatedTyreAgeAtLapStart",
    "IsPitOutLap",
    "PitLaneDuration_seconds",
    "StationaryStopDuration_seconds",
    "I1Speed_kmh",
    "I2Speed_kmh",
    "SpeedTrap_kmh",
    "LapStartUTC",
]

existing_preferred = [
    c
    for c in preferred_order
    if c in laps.columns
]

extra_cols = [
    c
    for c in laps.columns
    if (
        c not in existing_preferred
        and c not in {
            "meeting_key",
            "session_key",
            "segments_sector_1",
            "segments_sector_2",
            "segments_sector_3",
        }
    )
]

laps = laps[
    existing_preferred
    + extra_cols
]

sort_cols = [
    c
    for c in [
        "Driver",
        "DriverNumber",
        "LapNumber",
    ]
    if c in laps.columns
]

if sort_cols:
    laps = (
        laps
        .sort_values(sort_cols)
        .reset_index(drop=True)
    )


# ============================================================
# 10) COMPLETE-SECTOR SUBSET
# ============================================================

required_sector_cols = [
    "Sector1Time_seconds",
    "Sector2Time_seconds",
    "Sector3Time_seconds",
]

if all(
    c in laps.columns
    for c in required_sector_cols
):

    complete = laps.dropna(
        subset=required_sector_cols
    ).copy()

else:

    complete = laps.copy()


# ============================================================
# 11) SAVE EXCEL FILES
# ============================================================

master_path = (
    OUT_DIR
    / "japan_2026_race_all_laps_openf1.xlsx"
)

complete_path = (
    OUT_DIR
    / "japan_2026_race_complete_sector_times_openf1.xlsx"
)

save_formatted_excel(
    df=laps,
    path=master_path,
    sheet_name="All race laps",
    table_name="OpenF1AllRaceLaps",
)

save_formatted_excel(
    df=complete,
    path=complete_path,
    sheet_name="Complete sectors",
    table_name="OpenF1CompleteSectors",
)


# ============================================================
# 12) PRINT DIAGNOSTICS
# ============================================================

print("\n" + "=" * 72)
print("SUCCESS")
print("=" * 72)

print(
    f"Rows downloaded: {len(laps)}"
)

if "DriverNumber" in laps.columns:
    n_drivers = (
        laps["DriverNumber"]
        .nunique()
    )
else:
    n_drivers = "N/A"

print(
    f"Drivers: {n_drivers}"
)

print(
    "Complete S1/S2/S3 rows: "
    f"{len(complete)}"
)

for col in [
    "Sector1Time_seconds",
    "Sector2Time_seconds",
    "Sector3Time_seconds",
]:

    if col in laps.columns:

        print(
            f"Missing {col}: "
            f"{laps[col].isna().sum()}"
        )


if (
    "SectorSumMinusLap_seconds"
    in laps.columns
):

    diff = (
        laps[
            "SectorSumMinusLap_seconds"
        ]
        .dropna()
        .abs()
    )

    if len(diff) > 0:

        print(
            "Median "
            "|S1+S2+S3-LapTime|: "
            f"{diff.median():.6f} s"
        )

        print(
            "Max    "
            "|S1+S2+S3-LapTime|: "
            f"{diff.max():.6f} s"
        )


print("\nSaved:")

print(
    master_path.resolve()
)

print(
    complete_path.resolve()
)


preview = [
    c
    for c in [
        "Driver",
        "DriverNumber",
        "LapNumber",
        "LapTime_seconds",
        "Sector1Time_seconds",
        "Sector2Time_seconds",
        "Sector3Time_seconds",
        "Compound",
    ]
    if c in laps.columns
]

print("\nPreview:")

print(
    laps[preview]
    .head(25)
    .to_string(index=False)
)
